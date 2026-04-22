"""Generate Apple-style PPTX with refined TDM concept."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

FIGURES = "/vllm-workspace/lzn-pro/figures"
OUTPUT = "/vllm-workspace/lzn-pro/NPU_PD_TDM_Report.pptx"

# Apple palette
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFB, 0xFB, 0xFD)
BLACK = RGBColor(0x1D, 0x1D, 0x1F)
GRAY = RGBColor(0x86, 0x86, 0x8B)
LIGHT = RGBColor(0xD2, 0xD2, 0xD7)
BLUE = RGBColor(0x00, 0x71, 0xE3)
GREEN = RGBColor(0x34, 0xC7, 0x59)
RED = RGBColor(0xFF, 0x3B, 0x30)
ORANGE = RGBColor(0xFF, 0x9F, 0x0A)
PURPLE = RGBColor(0xAF, 0x52, 0xDE)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

def black_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BLACK

def tb(slide, l, t, w, h):
    return slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))

def text(tf, s, sz=18, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = s
    p.font.size = Pt(sz)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    return p

def para(tf, s, sz=18, color=BLACK, bold=False, align=PP_ALIGN.LEFT, before=6):
    p = tf.add_paragraph()
    p.text = s
    p.font.size = Pt(sz)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(2)
    return p

def thin_line(slide, y, x=1.5, w=10.3):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Pt(0.75))
    ln.fill.solid()
    ln.fill.fore_color.rgb = LIGHT
    ln.line.fill.background()

def tag(slide, x, y, label, color=BLUE):
    t = tb(slide, x, y, 3, 0.4)
    text(t.text_frame, label, sz=14, color=color, bold=True)

def fig(name):
    return os.path.join(FIGURES, name)


# ════════════════════════════════════════════════════════════
# 1. Title
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
black_bg(slide)

t = tb(slide, 1.5, 2.0, 10.3, 1.2)
text(t.text_frame, "NPU 单卡 PD 时分复用调度系统", sz=48, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

t2 = tb(slide, 1.5, 3.5, 10.3, 0.6)
text(t2.text_frame, "Single-NPU Prefill/Decode Time-Division Multiplexing",
     sz=22, color=GRAY, align=PP_ALIGN.CENTER)

t3 = tb(slide, 1.5, 5.5, 10.3, 0.5)
text(t3.text_frame, "2026-04-03    Qwen3-4B    2 \u00d7 Ascend 910B3",
     sz=16, color=GRAY, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 2. Problem
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "问题")

t2 = tb(slide, 1.5, 1.3, 10.3, 2.0)
text(t2.text_frame, "Prefill 与 Decode\n资源互补，却串行执行。",
     sz=44, color=BLACK, bold=True)

t3 = tb(slide, 1.5, 3.8, 10.3, 1.5)
text(t3.text_frame,
     "Prefill 是算力密集（AICore 69.8%），Decode 是带宽密集（HBM BW 25.1%）。\n"
     "单卡上两者交替空闲 —— 中低并发场景下，大量算力被浪费。",
     sz=20, color=GRAY)


# ════════════════════════════════════════════════════════════
# 3. Data: fig1 + fig2
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "实测数据")

slide.shapes.add_picture(fig("fig1_exp_b_resource_profile.png"),
                         Inches(0.8), Inches(1.4), Inches(5.8), Inches(5.2))
slide.shapes.add_picture(fig("fig2_exp_d_concurrency.png"),
                         Inches(6.8), Inches(1.4), Inches(5.8), Inches(5.2))


# ════════════════════════════════════════════════════════════
# 4. TDM Concept — the key new slide
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "核心概念")

t2 = tb(slide, 1.5, 1.2, 10.3, 1.0)
text(t2.text_frame, "TDM 不是每次做 P/D 二选一，\n而是控制时间窗口内的执行比例。",
     sz=38, color=BLACK, bold=True)

# Timeline visualization
t3 = tb(slide, 1.5, 2.8, 10.3, 0.4)
text(t3.text_frame, "时间轴  （每个字母 = 1 iteration = 1 batch = 1 ACL Graph replay）",
     sz=15, color=GRAY)

# Draw iteration blocks as colored rectangles
colors_map = {
    'D': BLUE, 'P': ORANGE, 'M': PURPLE
}
labels_top = list("DDDPDDDDPDDPPDDDDDDDPDD")
bw = 0.42  # block width
bh = 0.55
y_top = 3.4
x_start = 0.6
for i, ch in enumerate(labels_top):
    x = x_start + i * (bw + 0.04)
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y_top), Inches(bw), Inches(bh))
    rect.fill.solid()
    rect.fill.fore_color.rgb = colors_map[ch]
    rect.line.fill.background()
    tf = rect.text_frame
    tf.word_wrap = False
    text(tf, ch, sz=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# Annotation under blocks
t_ann = tb(slide, 0.6, 4.1, 4.5, 0.4)
text(t_ann.text_frame, "低负载：P 稀疏插入", sz=14, color=GRAY)
t_ann2 = tb(slide, 5.5, 4.1, 3.5, 0.4)
text(t_ann2.text_frame, "突发请求：P 密集", sz=14, color=ORANGE, bold=True)
t_ann3 = tb(slide, 9.0, 4.1, 3.0, 0.4)
text(t_ann3.text_frame, "恢复：P 稀疏", sz=14, color=GRAY)

# Second row: high concurrency with Mixed
labels_bot = list("DDMMDDDMDDD")
y_bot = 4.8
t_hc = tb(slide, 0.6, y_bot - 0.4, 5, 0.35)
text(t_hc.text_frame, "高并发场景退化为 Mixed（Chunked Prefill）", sz=15, color=GRAY)
for i, ch in enumerate(labels_bot):
    x = x_start + i * (bw + 0.04)
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y_bot), Inches(bw), Inches(bh))
    rect.fill.solid()
    rect.fill.fore_color.rgb = colors_map[ch]
    rect.line.fill.background()
    tf = rect.text_frame
    tf.word_wrap = False
    text(tf, ch, sz=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# Legend
leg_y = 5.7
for i, (ch, label, clr) in enumerate([
    ("D", "DECODE  纯 decode", BLUE),
    ("P", "PREFILL  纯 prefill", ORANGE),
    ("M", "MIXED  Chunked Prefill", PURPLE),
]):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(1.5 + i * 3.8), Inches(leg_y), Inches(0.4), Inches(0.35))
    rect.fill.solid()
    rect.fill.fore_color.rgb = clr
    rect.line.fill.background()
    tf = rect.text_frame
    text(tf, ch, sz=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    t_leg = tb(slide, 2.0 + i * 3.8, leg_y, 3.0, 0.35)
    text(t_leg.text_frame, label, sz=15, color=BLACK)

# Bottom: CP is a special case
t_bottom = tb(slide, 1.5, 6.4, 10.3, 0.6)
text(t_bottom.text_frame,
     "Chunked Prefill 是 TDM 的特例（全部 iteration 选 MIXED）。\nTDM 根据负载动态选择最优模式组合。",
     sz=18, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 5. Three modes — why we need all three
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "三种执行模式")

t2 = tb(slide, 1.5, 1.2, 10.3, 0.8)
text(t2.text_frame, "不是「分离一定好」或「混合一定好」",
     sz=36, color=BLACK, bold=True)

# Three columns
modes = [
    ("PREFILL", ORANGE, "纯 prefill iteration", "队列深\ndecode 空闲\n低并发"),
    ("DECODE", BLUE, "纯 decode iteration", "TPOT 紧张\n无新请求"),
    ("MIXED", PURPLE, "Chunked Prefill\nP+D 混合 batch", "高并发\n负载平稳\n混合效率更高"),
]
for i, (name, clr, desc, when) in enumerate(modes):
    x = 1.0 + i * 4.0
    # Mode name
    t_n = tb(slide, x, 2.5, 3.5, 0.6)
    text(t_n.text_frame, name, sz=32, color=clr, bold=True, align=PP_ALIGN.CENTER)
    # Description
    t_d = tb(slide, x, 3.3, 3.5, 0.7)
    text(t_d.text_frame, desc, sz=16, color=GRAY, align=PP_ALIGN.CENTER)
    # When
    t_w = tb(slide, x, 4.3, 3.5, 1.5)
    text(t_w.text_frame, when, sz=18, color=BLACK, align=PP_ALIGN.CENTER)

# Exp E data
thin_line(slide, 5.9)
t_exp = tb(slide, 1.5, 6.1, 10.3, 0.8)
text(t_exp.text_frame,
     "Exp E 实证：Phased 低并发 +10%，高并发 -20%  \u2192  需要动态选择，不能一刀切",
     sz=18, color=GRAY, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 6. Architecture overview
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.4, "三层架构")

t2 = tb(slide, 1.5, 0.9, 10.3, 0.6)
text(t2.text_frame, "Layer 1 决定模式（时间维度）  Layer 2 决定大小（容量维度）",
     sz=24, color=BLACK, bold=True)

slide.shapes.add_picture(fig("fig8_architecture_overview.png"),
                         Inches(1.8), Inches(1.8), Inches(9.7), Inches(5.2))


# ════════════════════════════════════════════════════════════
# 7. Why iteration-level
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "为什么是 iteration 级")

t2 = tb(slide, 1.5, 1.3, 10.3, 1.0)
text(t2.text_frame, "ACL Graph 决定了 iteration\n是 NPU 上最小的原子执行单元。",
     sz=36, color=BLACK, bold=True)

# Two columns
left = tb(slide, 1.5, 3.2, 5.0, 3.0)
tf = left.text_frame
text(tf, "算子级交错", sz=22, color=RED, bold=True)
para(tf, "FULL 模式 graph 不可拆分", sz=18, color=GRAY, before=12)
para(tf, "PIECEWISE 劣化 26%", sz=18, color=GRAY, before=8)
para(tf, "Stream 硬上限 2048", sz=18, color=GRAY, before=8)
para(tf, "不可行", sz=20, color=RED, bold=True, before=16)

right = tb(slide, 7.3, 3.2, 5.0, 3.0)
tf = right.text_frame
text(tf, "Iteration 级时分", sz=22, color=GREEN, bold=True)
para(tf, "纯调度层改动", sz=18, color=GRAY, before=12)
para(tf, "Graph 自然执行边界", sz=18, color=GRAY, before=8)
para(tf, "切换零开销", sz=18, color=GRAY, before=8)
para(tf, "可行，改动最小", sz=20, color=GREEN, bold=True, before=16)

# Equivalence
t_eq = tb(slide, 1.5, 6.3, 10.3, 0.5)
text(t_eq.text_frame,
     "1 iteration  =  1 batch  =  1 ACL Graph replay    "
     "Iteration 是最小切换单元，调度策略看窗口级比例",
     sz=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 8. Phase Ratio Controller — V1
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "Layer 1  Phase Ratio Controller")

t2 = tb(slide, 1.5, 1.2, 10.3, 0.8)
text(t2.text_frame, "V1  优先级规则 + 防饿死", sz=36, color=BLACK, bold=True)

body = tb(slide, 1.5, 2.4, 10.3, 4.5)
tf = body.text_frame
tf.word_wrap = True
text(tf, "", sz=4, color=BG)
rules = [
    ("TPOT 即将违约", "DECODE", "最高优先级，保 decode SLO"),
    ("TTFT 即将违约", "PREFILL", "防首 token 超时"),
    ("连续 decode > 上限", "PREFILL", "防 prefill 饿死"),
    ("有等待 & 无 decode", "PREFILL", "空闲时做 prefill"),
    ("高并发 & 有等待", "MIXED", "混合执行效率更高"),
    ("无等待请求", "DECODE", "全力完成已有请求"),
]
for cond, act, why in rules:
    clr = BLUE if act == "DECODE" else (ORANGE if act == "PREFILL" else PURPLE)
    para(tf, f"{cond}    \u2192    {act}", sz=20, color=BLACK, bold=True, before=14)
    para(tf, why, sz=16, color=GRAY, before=2)

# Limitation note
t_lim = tb(slide, 1.5, 6.5, 10.3, 0.5)
text(t_lim.text_frame, "局限：本质 reactive，等 SLO 快违约才切换，可能震荡",
     sz=16, color=RED, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 9. Phase Ratio Controller — V2 AIMD
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "Layer 1  Phase Ratio Controller")

t2 = tb(slide, 1.5, 1.2, 10.3, 0.8)
text(t2.text_frame, "V2  AIMD 自适应比例控制", sz=36, color=BLACK, bold=True)

t3 = tb(slide, 1.5, 2.3, 10.3, 0.8)
text(t3.text_frame,
     "核心变量：prefill 插入率 r\u209a\n"
     "不是每次做决策，而是维护一个平滑变化的比例。",
     sz=20, color=GRAY)

body = tb(slide, 1.5, 3.5, 10.3, 2.5)
tf = body.text_frame
tf.word_wrap = True
text(tf, "", sz=4, color=BG)
para(tf, "TPOT 违约（拥塞信号）  \u2192  r\u209a \u00d7 0.5    Multiplicative Decrease",
     sz=22, color=RED, bold=True, before=16)
para(tf, "TPOT 达标 & Q > 0     \u2192  r\u209a + \u03b1      Additive Increase",
     sz=22, color=GREEN, bold=True, before=16)
para(tf, "高并发 R > threshold   \u2192  切换 MIXED 模式",
     sz=22, color=PURPLE, bold=True, before=16)

# Anti-starvation
thin_line(slide, 5.4)
t_as = tb(slide, 1.5, 5.7, 10.3, 1.2)
tf = t_as.text_frame
tf.word_wrap = True
text(tf, "防饿死保证", sz=20, color=BLACK, bold=True)
para(tf, "r\u209a 有下界（Q > 0 时保证 prefill 不被无限延迟）  "
     "  连续 decode 有上界 max_consecutive",
     sz=17, color=GRAY, before=8)
para(tf, "渐进收敛到当前负载的最优 P:D:M 比例，无需手调阈值",
     sz=17, color=BLUE, bold=True, before=8)


# ════════════════════════════════════════════════════════════
# 10. Comparison — fig6 + fig7
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "2 卡公平对比")

t2 = tb(slide, 1.5, 1.1, 10.3, 0.6)
text(t2.text_frame, "四种方案实测", sz=32, color=BLACK, bold=True)

slide.shapes.add_picture(fig("fig6_exp_e3_fair_compare.png"),
                         Inches(0.8), Inches(2.0), Inches(5.8), Inches(4.8))
slide.shapes.add_picture(fig("fig7_exp_e3_latency.png"),
                         Inches(6.8), Inches(2.0), Inches(5.8), Inches(4.8))


# ════════════════════════════════════════════════════════════
# 11. Comparison stats
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)

items = [
    ("Unified + CP", "+5 ~ 11%", "强 baseline", BLUE),
    ("Phased TP=2", "+2 ~ 7%", "一刀切，非动态", ORANGE),
    ("Disagg 1P1D", "-2 ~ 12%", "物理分离，浪费", RED),
]
x = 1.0
for name, pct, note, clr in items:
    t_n = tb(slide, x, 1.5, 3.5, 0.5)
    text(t_n.text_frame, name, sz=16, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
    t_p = tb(slide, x, 2.2, 3.5, 1.2)
    text(t_p.text_frame, pct, sz=52, color=clr, bold=True, align=PP_ALIGN.CENTER)
    t_no = tb(slide, x, 3.6, 3.5, 0.5)
    text(t_no.text_frame, note, sz=16, color=GRAY, align=PP_ALIGN.CENTER)
    x += 4.0

t4 = tb(slide, 1.5, 5.0, 10.3, 1.2)
text(t4.text_frame,
     "TDM 目标：在 Phased 验证的「分离有收益」基础上，\n"
     "通过 AIMD 动态控制 P:D:Mixed 比例，超越 Chunked Prefill。",
     sz=20, color=BLACK, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 12. Graph-aware problem — fig3 + fig4
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "Layer 2  Graph 感知 Batch 组织")

t2 = tb(slide, 1.5, 1.3, 10.3, 1.0)
text(t2.text_frame, "48 种编译 shape，\npadding 浪费最高 62.5%。",
     sz=40, color=BLACK, bold=True)

slide.shapes.add_picture(fig("fig3_exp_a_acl_graph.png"),
                         Inches(0.8), Inches(3.2), Inches(5.8), Inches(3.8))
slide.shapes.add_picture(fig("fig4_exp_f_graph_aware.png"),
                         Inches(6.8), Inches(3.2), Inches(5.8), Inches(3.8))


# ════════════════════════════════════════════════════════════
# 13. Layer 1 + Layer 2 relationship
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "Layer 1 + Layer 2 协作")

t2 = tb(slide, 1.5, 1.2, 10.3, 0.7)
text(t2.text_frame, "串行决策 + 反馈修正", sz=32, color=BLACK, bold=True)

steps = [
    ("01", "Layer 1 决定模式", "PREFILL / DECODE / MIXED   （时间维度：做什么）", BLUE),
    ("02", "Layer 2 决定 Batch Size", "对齐最近的编译 shape   （容量维度：做多大）", GREEN),
    ("03", "Layer 2 反馈 Layer 1", "凑不够好 shape \u2192 建议延迟切换", ORANGE),
]
y = 2.3
for num, title, sub, clr in steps:
    t_num = tb(slide, 1.5, y, 1.0, 0.7)
    text(t_num.text_frame, num, sz=44, color=clr, bold=True)
    t_title = tb(slide, 2.8, y, 9.0, 0.4)
    text(t_title.text_frame, title, sz=24, color=BLACK, bold=True)
    t_sub = tb(slide, 2.8, y + 0.5, 9.0, 0.4)
    text(t_sub.text_frame, sub, sz=17, color=GRAY)
    y += 1.4

# Example
thin_line(slide, 5.8)
t_ex = tb(slide, 1.5, 6.0, 10.3, 1.0)
tf = t_ex.text_frame
tf.word_wrap = True
text(tf, "示例：running=33，新来 2 个请求", sz=18, color=BLACK, bold=True)
para(tf, "Layer 1: DECODE (ttft 充裕)   Layer 2: 33\u219240 浪费 17.5%, 建议等降到 32   "
     "最终: 保持 DECODE, 32 后切 PREFILL",
     sz=16, color=GRAY, before=6)


# ════════════════════════════════════════════════════════════
# 14. Originality
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "原创性")

t2 = tb(slide, 1.5, 1.8, 10.3, 2.5)
text(t2.text_frame,
     "调研 17 篇论文，\n无一将编译图约束\n引入调度决策。",
     sz=44, color=BLACK, bold=True)

t3 = tb(slide, 1.5, 4.8, 10.3, 1.0)
text(t3.text_frame,
     "12 篇核心 PD 论文 + 5 篇 TDM 论文\n"
     "（FaST-GShare / GaiaGPU / TGS / FaaSwap / Dilu）",
     sz=18, color=GRAY)

t4 = tb(slide, 1.5, 6.0, 10.3, 0.6)
text(t4.text_frame,
     "TDM 作为 CP 超集 + Graph-Aware Batch Shaping + NPU 首创",
     sz=18, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 15. Challenges
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "挑战")

t2 = tb(slide, 1.5, 1.3, 10.3, 0.8)
text(t2.text_frame, "三个待解决的问题", sz=36, color=BLACK, bold=True)

challenges = [
    ("核心模块尚未实现",
     "Phase Ratio Controller V1 实现 + 阈值确定\nGraph-Aware Batch Shaping 集成 + Mixed 模式支持"),
    ("小模型天花板低",
     "4B 模型 prefill 仅 30ms，干扰微弱\n30B+ 才是真正目标场景"),
    ("Chunked Prefill 是强 baseline",
     "CP 已 +5~11%，TDM 需论证增量价值\n但 NPU 上 CP 受限，TDM 作为超集价值可能更大"),
]
y = 2.8
for title, desc in challenges:
    t_t = tb(slide, 1.5, y, 10.3, 0.4)
    text(t_t.text_frame, title, sz=24, color=BLACK, bold=True)
    t_d = tb(slide, 1.5, y + 0.5, 10.3, 0.8)
    text(t_d.text_frame, desc, sz=18, color=GRAY)
    y += 1.5


# ════════════════════════════════════════════════════════════
# 16. Next steps
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
tag(slide, 1.5, 0.6, "计划")

t2 = tb(slide, 1.5, 1.3, 10.3, 0.8)
text(t2.text_frame, "下一步", sz=36, color=BLACK, bold=True)

plans = [
    ("1 ~ 2 周", BLUE, [
        "Phase Ratio Controller V1 原型",
        "End-to-end benchmark vs Unified+CP",
        "离线 latency 画像表采集",
    ]),
    ("2 ~ 4 周", ORANGE, [
        "V2 AIMD 自适应 + Mixed 模式",
        "Graph-Aware Batch Shaping 集成",
        "Qwen3-30B 大模型测试",
    ]),
    ("长期", GRAY, [
        "论文撰写",
        "向 vllm-ascend 社区提交 PR",
    ]),
]
x = 1.0
for period, clr, items in plans:
    t_p = tb(slide, x, 2.6, 3.5, 0.5)
    text(t_p.text_frame, period, sz=26, color=clr, bold=True, align=PP_ALIGN.CENTER)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x + 0.5), Inches(3.2), Inches(2.5), Pt(3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = clr
    bar.line.fill.background()
    t_items = tb(slide, x + 0.2, 3.6, 3.2, 3.0)
    tf = t_items.text_frame
    tf.word_wrap = True
    text(tf, "", sz=4, color=BG)
    for item in items:
        para(tf, item, sz=18, color=BLACK, before=14)
    x += 4.0


# ════════════════════════════════════════════════════════════
# 17. End
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
black_bg(slide)

t = tb(slide, 1.5, 2.5, 10.3, 1.5)
text(t.text_frame, "谢谢", sz=56, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

t2 = tb(slide, 1.5, 4.2, 10.3, 0.6)
text(t2.text_frame, "Q & A", sz=24, color=GRAY, align=PP_ALIGN.CENTER)


# ── Save ─────────────────────────────────────────────────
prs.save(OUTPUT)
print(f"Saved: {OUTPUT}")
print(f"Slides: {len(prs.slides)}")
